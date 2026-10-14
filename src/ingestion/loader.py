import os
import re
from pathlib import Path
from typing import List, Optional, Union
from langchain_core.documents import Document
from pypdf import PdfReader

from src.ingestion.vision import VisionDescriber


class DocumentLoader:
    """Document loader supporting PDF, PPTX, PPT, TXT, Markdown, and slide images with a uniform metadata schema."""

    SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".png", ".jpg", ".jpeg", ".pptx", ".ppt"}
    IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}
    SLIDE_EXTENSIONS = {".pptx", ".ppt"}

    def __init__(
        self,
        vision_describer: Optional[VisionDescriber] = None,
        extract_visuals: bool = True,
    ):
        self.extract_visuals = extract_visuals
        if vision_describer is not None:
            self.vision_describer = vision_describer
        elif extract_visuals:
            try:
                self.vision_describer = VisionDescriber()
            except Exception:
                self.vision_describer = None
        else:
            self.vision_describer = None

    def load(self, file_path: Union[str, Path]) -> List[Document]:
        path = Path(file_path).resolve()

        if not path.is_file():
            raise FileNotFoundError(f"File not found: {path}")

        ext = path.suffix.lower()
        if ext not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file format: '{ext}'. Supported formats: {sorted(self.SUPPORTED_EXTENSIONS)}"
            )

        if ext == ".pdf":
            return self._load_pdf(path)
        if ext == ".pptx":
            return self._load_pptx(path)
        if ext == ".ppt":
            return self._load_ppt(path)
        if ext in self.IMAGE_EXTENSIONS:
            return self._load_image(path, ext)
        return self._load_text(path, ext)

    def _load_pdf(self, path: Path) -> List[Document]:
        documents: List[Document] = []
        reader = PdfReader(str(path))

        page_count = len(reader.pages)
        # A full textbook has hundreds of embedded figures. Describing each one
        # stalls the upload and then the vision API rejects the rest of the book.
        allow_vision = self.extract_visuals and self.vision_describer is not None and page_count <= 40
        vision_calls = 0

        for page_idx, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            diagram_descriptions: List[str] = []

            # Check if page is landscape oriented (characteristic of presentation slides)
            is_slide = False
            try:
                mb = page.mediabox
                if mb and mb.width > mb.height:
                    is_slide = True
            except Exception:
                pass

            # Detect short header title
            detected_title = ""
            if text:
                first_line = text.strip().split("\n")[0].strip()
                if 3 <= len(first_line) <= 120 and not first_line.startswith("http"):
                    detected_title = first_line

            if allow_vision and vision_calls < 8:
                try:
                    for img_obj in page.images:
                        img_bytes = img_obj.data
                        if self.vision_describer.is_meaningful_diagram(img_bytes):
                            desc = self.vision_describer.describe_diagram(
                                img_bytes,
                                context_hint=f"{path.name} (Page {page_idx})",
                            )
                            if desc and desc not in diagram_descriptions:
                                diagram_descriptions.append(desc)
                                vision_calls += 1
                            if vision_calls >= 8:
                                break
                except Exception:
                    # Robust error isolation: Never fail text extraction on image parsing hiccups
                    pass

            if diagram_descriptions:
                visual_blocks = "\n\n".join(
                    f"[Visual Diagram / Slide Figure on Page {page_idx}]:\n{d}"
                    for d in diagram_descriptions
                )
                page_content = f"{text}\n\n{visual_blocks}".strip() if text.strip() else visual_blocks
            else:
                page_content = text

            meta = {
                "source": path.name,
                "file_path": str(path),
                "file_type": "pdf",
                "page": int(page_idx),
            }
            if is_slide:
                meta["is_slide"] = True
            if detected_title:
                meta["slide_title"] = detected_title
            if diagram_descriptions:
                meta["has_visuals"] = True
                meta["visual_count"] = len(diagram_descriptions)

            documents.append(
                Document(
                    page_content=page_content,
                    metadata=meta,
                )
            )
        return documents

    def _load_pptx(self, path: Path) -> List[Document]:
        from pptx import Presentation

        documents: List[Document] = []
        prs = Presentation(str(path))
        slide_count = len(prs.slides)
        allow_vision = self.extract_visuals and self.vision_describer is not None and slide_count <= 40
        vision_calls = 0

        for slide_idx, slide in enumerate(prs.slides, start=1):
            title = ""
            if slide.shapes.title and slide.shapes.title.text:
                title = slide.shapes.title.text.strip()

            body_lines: List[str] = []
            table_lines: List[str] = []
            diagram_descriptions: List[str] = []

            for shape in slide.shapes:
                if shape.has_text_frame:
                    if slide.shapes.title and shape == slide.shapes.title:
                        continue
                    for paragraph in shape.text_frame.paragraphs:
                        p_text = paragraph.text.strip()
                        if p_text:
                            bullet = "- " if paragraph.level > 0 else ""
                            indent = "  " * paragraph.level
                            body_lines.append(f"{indent}{bullet}{p_text}")

                elif shape.has_table:
                    table = shape.table
                    rows: List[str] = []
                    for row in table.rows:
                        row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                        if any(row_cells):
                            rows.append("| " + " | ".join(row_cells) + " |")
                    if rows:
                        table_lines.append("\n".join(rows))

                elif hasattr(shape, "image") or getattr(shape, "shape_type", None) == 13:
                    if allow_vision and vision_calls < 8:
                        try:
                            image_blob = shape.image.blob
                            if self.vision_describer.is_meaningful_diagram(image_blob):
                                desc = self.vision_describer.describe_diagram(
                                    image_blob,
                                    context_hint=f"{path.name} (Slide {slide_idx})",
                                )
                                if desc and desc not in diagram_descriptions:
                                    diagram_descriptions.append(desc)
                                    vision_calls += 1
                                if vision_calls >= 8:
                                    break
                        except Exception:
                            pass

            notes_text = ""
            if getattr(slide, "has_notes_slide", False) and slide.notes_slide:
                if getattr(slide.notes_slide, "notes_text_frame", None):
                    notes_text = slide.notes_slide.notes_text_frame.text.strip()

            header = f"[Slide {slide_idx}: {title}]" if title else f"[Slide {slide_idx}]"
            blocks: List[str] = [header]
            if body_lines:
                blocks.append("\n".join(body_lines))
            if table_lines:
                blocks.append("[Slide Table]:\n" + "\n\n".join(table_lines))
            if notes_text:
                blocks.append(f"[Speaker Notes]:\n{notes_text}")
            if diagram_descriptions:
                blocks.append("\n\n".join(
                    f"[Slide Figure / Visual on Slide {slide_idx}]:\n{d}"
                    for d in diagram_descriptions
                ))

            page_content = "\n\n".join(blocks).strip()

            meta = {
                "source": path.name,
                "file_path": str(path),
                "file_type": "pptx",
                "page": int(slide_idx),
                "is_slide": True,
            }
            if title:
                meta["slide_title"] = title
            if diagram_descriptions:
                meta["has_visuals"] = True
                meta["visual_count"] = len(diagram_descriptions)

            documents.append(Document(page_content=page_content, metadata=meta))

        return documents

    def _load_ppt(self, path: Path) -> List[Document]:
        # Handle cases where .ppt is actually a renamed .pptx (zip archive)
        try:
            with open(path, "rb") as f:
                header = f.read(4)
            if header == b"PK\x03\x04":
                docs = self._load_pptx(path)
                for d in docs:
                    d.metadata["file_type"] = "ppt"
                return docs
        except Exception:
            pass

        # Parse binary PPT via olefile
        raw_text_parts: List[str] = []
        try:
            import olefile
            if olefile.isOleFile(str(path)):
                with olefile.OleFileIO(str(path)) as ole:
                    if ole.exists("PowerPoint Document"):
                        data = ole.openstream("PowerPoint Document").read()
                        # Extract ASCII printable sequences of length >= 4
                        ascii_strings = [
                            s.decode("latin-1", errors="ignore").strip()
                            for s in re.findall(rb"[\x20-\x7E\t\r\n]{4,}", data)
                        ]
                        # Extract UTF-16LE printable sequences
                        utf16_strings = []
                        for s in re.findall(rb"(?:[\x20-\x7E]\x00){4,}", data):
                            try:
                                decoded = s.decode("utf-16le", errors="ignore").strip()
                                if decoded:
                                    utf16_strings.append(decoded)
                            except Exception:
                                pass
                        all_strings = [s for s in ascii_strings + utf16_strings if len(s) > 3]
                        if all_strings:
                            raw_text_parts = all_strings
        except Exception:
            pass

        if not raw_text_parts:
            # Fallback to direct raw binary extraction
            try:
                data = path.read_bytes()
                strings = [
                    s.decode("latin-1", errors="ignore").strip()
                    for s in re.findall(rb"[\x20-\x7E\t\r\n]{4,}", data)
                ]
                raw_text_parts = [s for s in strings if len(s) > 3]
            except Exception:
                raw_text_parts = []

        combined_text = "\n".join(raw_text_parts).strip() if raw_text_parts else f"[Slide Presentation: {path.name}]"
        return [
            Document(
                page_content=combined_text,
                metadata={
                    "source": path.name,
                    "file_path": str(path),
                    "file_type": "ppt",
                    "page": 1,
                    "is_slide": True,
                },
            )
        ]

    def _load_image(self, path: Path, ext: str) -> List[Document]:
        data = path.read_bytes()
        desc = ""
        if self.vision_describer:
            desc = self.vision_describer.describe_diagram(
                data, context_hint=f"Lecture Slide Image: {path.name}"
            )
        if not desc:
            desc = f"Academic diagram / slide image: {path.name}"

        content = f"[Lecture Slide / Diagram: {path.name}]:\n{desc}"
        return [
            Document(
                page_content=content,
                metadata={
                    "source": path.name,
                    "file_path": str(path),
                    "file_type": "image",
                    "page": 1,
                    "has_visuals": True,
                    "visual_count": 1,
                },
            )
        ]

    def _load_text(self, path: Path, ext: str) -> List[Document]:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()
        return [
            Document(
                page_content=content,
                metadata={
                    "source": path.name,
                    "file_path": str(path),
                    "file_type": ext.lstrip("."),
                    "page": 1,
                },
            )
        ]
