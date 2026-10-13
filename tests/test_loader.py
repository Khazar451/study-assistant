from pathlib import Path
import pytest
from pypdf import PdfWriter
from src.ingestion.loader import DocumentLoader


def test_load_txt_file_uniform_metadata(tmp_path):
    loader = DocumentLoader()
    txt_file = tmp_path / "lecture_notes.txt"
    txt_file.write_text("Introduction to Calculus and limits.", encoding="utf-8")

    docs = loader.load(txt_file)

    assert len(docs) == 1
    assert docs[0].page_content == "Introduction to Calculus and limits."
    assert docs[0].metadata["source"] == "lecture_notes.txt"
    assert docs[0].metadata["file_path"] == str(txt_file.resolve())
    assert docs[0].metadata["file_type"] == "txt"
    assert docs[0].metadata["page"] == 1


def test_load_md_file_uniform_metadata(tmp_path):
    loader = DocumentLoader()
    md_file = tmp_path / "physics_summary.md"
    md_file.write_text("# Physics\nNewton's first law of motion.", encoding="utf-8")

    docs = loader.load(md_file)

    assert len(docs) == 1
    assert "Newton's first law" in docs[0].page_content
    assert docs[0].metadata["source"] == "physics_summary.md"
    assert docs[0].metadata["file_path"] == str(md_file.resolve())
    assert docs[0].metadata["file_type"] == "md"
    assert docs[0].metadata["page"] == 1


def test_load_unsupported_extension_raises(tmp_path):
    loader = DocumentLoader()
    bad_file = tmp_path / "script.exe"
    bad_file.write_text("binary content", encoding="utf-8")

    with pytest.raises(ValueError, match="Unsupported file format: '.exe'"):
        loader.load(bad_file)


def test_load_missing_file_raises(tmp_path):
    loader = DocumentLoader()
    missing_file = tmp_path / "non_existent.pdf"

    with pytest.raises(FileNotFoundError, match="File not found"):
        loader.load(missing_file)


def test_load_pdf_file_multipage(tmp_path):
    loader = DocumentLoader()
    pdf_path = tmp_path / "sample_multipage.pdf"

    # Create a 2-page PDF with pypdf
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.add_blank_page(width=72, height=72)
    with open(pdf_path, "wb") as f:
        writer.write(f)

    docs = loader.load(pdf_path)

    assert len(docs) == 2
    assert docs[0].metadata["page"] == 1
    assert docs[1].metadata["page"] == 2
    assert docs[0].metadata["file_type"] == "pdf"
    assert docs[0].metadata["source"] == "sample_multipage.pdf"
    assert docs[0].metadata["file_path"] == str(pdf_path.resolve())


def test_load_png_image(tmp_path):
    from PIL import Image
    from unittest.mock import MagicMock
    from src.ingestion.vision import VisionDescriber

    img_path = tmp_path / "system_architecture.png"
    Image.new("RGB", (300, 300), color="blue").save(img_path)

    mock_vision = MagicMock(spec=VisionDescriber)
    mock_vision.describe_diagram.return_value = "System Architecture: Microservices connected via Kafka."

    loader = DocumentLoader(vision_describer=mock_vision)
    docs = loader.load(img_path)

    assert len(docs) == 1
    assert "System Architecture" in docs[0].page_content
    assert docs[0].metadata["source"] == "system_architecture.png"
    assert docs[0].metadata["file_type"] == "image"
    assert docs[0].metadata["page"] == 1
    assert docs[0].metadata["has_visuals"] is True
    assert docs[0].metadata["visual_count"] == 1


def test_load_pdf_with_visual_diagram(tmp_path):
    from unittest.mock import MagicMock
    from src.ingestion.vision import VisionDescriber

    pdf_path = tmp_path / "lecture_with_slide.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with open(pdf_path, "wb") as f:
        writer.write(f)

    mock_vision = MagicMock(spec=VisionDescriber)
    mock_vision.is_meaningful_diagram.return_value = True
    mock_vision.describe_diagram.return_value = "Confusion Matrix with precision 94%."

    loader = DocumentLoader(vision_describer=mock_vision, extract_visuals=True)
    docs = loader.load(pdf_path)

    assert len(docs) == 1
    assert docs[0].metadata["file_type"] == "pdf"


def test_load_pptx_slides(tmp_path):
    import pptx

    pptx_path = tmp_path / "quantum_mechanics.pptx"
    prs = pptx.Presentation()

    # Slide 1: Title
    slide_layout = prs.slide_layouts[0]
    slide1 = prs.slides.add_slide(slide_layout)
    slide1.shapes.title.text = "Quantum Mechanics Basics"
    slide1.placeholders[1].text = "Wave-particle duality and operators"

    # Slide 2: Bullets & Speaker Notes
    slide_layout2 = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(slide_layout2)
    slide2.shapes.title.text = "Schrodinger Equation"
    tf = slide2.placeholders[1].text_frame
    tf.text = "Time-dependent equation"
    p2 = tf.add_paragraph()
    p2.text = "Hamiltonian operator acts on state vector"
    p2.level = 1

    slide2.notes_slide.notes_text_frame.text = "Exam Tip: Memorize the 1D wave equation solution."

    # Slide 3: Table
    slide_layout3 = prs.slide_layouts[6]  # Blank
    slide3 = prs.slides.add_slide(slide_layout3)
    table_shape = slide3.shapes.add_table(rows=2, cols=2, left=100, top=100, width=400, height=200)
    table = table_shape.table
    table.cell(0, 0).text = "Particle"
    table.cell(0, 1).text = "Spin"
    table.cell(1, 0).text = "Electron"
    table.cell(1, 1).text = "1/2"

    prs.save(str(pptx_path))

    loader = DocumentLoader()
    docs = loader.load(pptx_path)

    assert len(docs) == 3

    # Check Slide 1
    assert docs[0].metadata["source"] == "quantum_mechanics.pptx"
    assert docs[0].metadata["file_type"] == "pptx"
    assert docs[0].metadata["page"] == 1
    assert docs[0].metadata["is_slide"] is True
    assert docs[0].metadata["slide_title"] == "Quantum Mechanics Basics"
    assert "Wave-particle duality" in docs[0].page_content

    # Check Slide 2 (Bullets + Speaker Notes)
    assert docs[1].metadata["page"] == 2
    assert docs[1].metadata["slide_title"] == "Schrodinger Equation"
    assert "Time-dependent equation" in docs[1].page_content
    assert "Hamiltonian operator" in docs[1].page_content
    assert "[Speaker Notes]:" in docs[1].page_content
    assert "Exam Tip: Memorize" in docs[1].page_content

    # Check Slide 3 (Table)
    assert docs[2].metadata["page"] == 3
    assert "[Slide Table]:" in docs[2].page_content
    assert "Electron" in docs[2].page_content
    assert "1/2" in docs[2].page_content


def test_load_landscape_pdf_slide_detection(tmp_path):
    loader = DocumentLoader()
    pdf_path = tmp_path / "presentation_slides.pdf"

    writer = PdfWriter()
    # Landscape orientation: width (192) > height (108) (16:9 ratio)
    page = writer.add_blank_page(width=192, height=108)
    with open(pdf_path, "wb") as f:
        writer.write(f)

    docs = loader.load(pdf_path)

    assert len(docs) == 1
    assert docs[0].metadata["file_type"] == "pdf"
    assert docs[0].metadata["page"] == 1
    assert docs[0].metadata["is_slide"] is True


def test_load_ppt_file(tmp_path):
    loader = DocumentLoader()
    ppt_path = tmp_path / "legacy_lecture.ppt"
    ppt_path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 500)

    docs = loader.load(ppt_path)

    assert len(docs) == 1
    assert docs[0].metadata["source"] == "legacy_lecture.ppt"
    assert docs[0].metadata["file_type"] == "ppt"
    assert docs[0].metadata["is_slide"] is True


